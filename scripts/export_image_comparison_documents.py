#!/usr/bin/env python3
"""Export the image comparison into Markdown and Word documents."""

import argparse
import csv
import hashlib
import json
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageOps


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_COMPARISON_DIR = ROOT / "MMKGC/outputs/mkgw_image_comparison_50"
DEFAULT_DESCRIPTIONS = ROOT / "MMKGC/outputs/mkgw_image_test_50_20261014_fixed2/entity_descriptions.csv"


def load_descriptions(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return {row["entity_key"]: row for row in csv.DictReader(handle)}


def relpath(path, root):
    return Path(path).resolve().relative_to(root.resolve()).as_posix()


def md_escape(value):
    return str(value or "").replace("|", "\\|")


def format_score(value):
    if value is None:
        return "N/A"
    return f"{float(value):.6f}"


def load_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def load_generation_prompts(entity):
    """Load the exact prompts used for the two displayed best images."""
    mmkgc_candidate = entity["mmkgc"]["candidate_id"]
    mmkgc_source = Path(entity["mmkgc"]["source"])
    mmkgc_metadata_path = mmkgc_source.with_suffix(".json")
    mmkgc_metadata = load_json(mmkgc_metadata_path)
    mmkgc_prompt = mmkgc_metadata.get("positive_prompt", "")

    legacy_source = Path(entity["legacy_knowledge_bridger"]["source"])
    legacy_dir = legacy_source.parent
    legacy_prompts = load_json(legacy_dir / "prompts.json")
    legacy_index = int(entity["legacy_knowledge_bridger"]["candidate_index"])
    legacy_prompt = legacy_prompts[legacy_index] if legacy_index < len(legacy_prompts) else ""

    return {
        "mmkgc_candidate": mmkgc_candidate,
        "mmkgc": mmkgc_prompt,
        "legacy_candidate": f"id_{legacy_index}.jpg",
        "legacy": legacy_prompt,
    }


def prepare_word_image(source, cache_dir):
    """Normalize images to conservative RGB JPEGs for Word compatibility."""
    digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:12]
    target = cache_dir / f"{source.stem}_{digest}.jpg"
    if target.exists():
        return target
    with Image.open(source) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((1200, 1200), Image.Resampling.LANCZOS)
        image.save(target, format="JPEG", quality=88, optimize=True, progressive=False)
    return target


def image_width(path, max_width=2.1):
    with Image.open(path) as image:
        return Inches(max_width)


def add_image_cell(cell, title, image_path, detail):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(10)
    paragraph.add_run("\n")
    width = image_width(image_path)
    image_run = paragraph.add_run()
    image_run.add_picture(str(image_path), width=width)
    detail_paragraph = cell.add_paragraph(detail)
    detail_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in detail_paragraph.runs:
        run.font.size = Pt(8)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--comparison-dir", type=Path, default=DEFAULT_COMPARISON_DIR)
    parser.add_argument("--descriptions", type=Path, default=DEFAULT_DESCRIPTIONS)
    args = parser.parse_args()

    comparison_dir = args.comparison_dir.resolve()
    manifest = json.loads((comparison_dir / "comparison_manifest.json").read_text(encoding="utf-8"))
    descriptions = load_descriptions(args.descriptions)
    legacy_summary = manifest["legacy_rescore_summary"]
    legacy_averages = legacy_summary["averages"]
    score_note_text = (
        "两种方法的候选图片均使用 MMKGC 的本地 CLIP ViT-B/32 和 "
        "BLIP/ITM 评分器计算图文匹配度，但各自对照本方法生成图片时使用的文本。"
        "CLIP 与 BLIP 原始 0~100 分数除以 100 后，按 "
        "(0.2×CLIP + 0.3×BLIP) / 0.5 = 0.4×CLIP + 0.6×BLIP 融合。"
        "legacy 的 250 张候选已全部重算，并按该融合分重新选择每个实体的最高候选；"
        "其全部候选平均融合分为 "
        f"{format_score(legacy_averages['clip_blip_score'])}。"
        "该分数衡量候选图片与各自生成文本的匹配度，不是相同参考文本下的绝对图片质量。"
    )

    md_parts = [
        "# MMKGC 与 legacy Knowledge-Bridger 图片对比",
        "",
        "共 50 个实体。每个实体包含文本描述、原图、MMKGC 最高分生成图和 legacy Knowledge-Bridger 最高分生成图。",
        "",
        f"> 分数说明：{score_note_text}",
        "",
        "## 实体目录",
        "",
    ]
    for entity in manifest["entities"]:
        anchor = f"entity-{entity['idx']:03d}"
        md_parts.append(f"- [{entity['idx'] + 1}. {md_escape(entity['label'])}](#{anchor})（{entity['entity_key']}）")
    md_parts.append("")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    document.add_heading("MMKGC 与 legacy Knowledge-Bridger 图片对比", level=0)
    document.add_paragraph(
        "共 50 个实体。每个实体包含文本描述、原图、MMKGC 最高分生成图和 "
        "legacy Knowledge-Bridger 最高分生成图。"
    )
    note = document.add_paragraph()
    note.add_run("注意：").bold = True
    note.add_run(score_note_text)
    document.add_heading("实体目录", level=1)
    for entity in manifest["entities"]:
        document.add_paragraph(f"{entity['idx'] + 1}. {entity['label']} ({entity['entity_key']})", style="List Bullet")

    with tempfile.TemporaryDirectory(prefix="comparison-word-images-") as temp_dir:
        cache_dir = Path(temp_dir)
        for entity in manifest["entities"]:
            idx = entity["idx"]
            key = entity["entity_key"]
            label = entity["label"]
            row = descriptions.get(key, {})
            description = row.get("description", "")
            prompts = load_generation_prompts(entity)
            original = comparison_dir / relpath(entity["original"]["copied"], comparison_dir)
            mmkgc = comparison_dir / relpath(entity["mmkgc"]["copied"], comparison_dir)
            legacy = comparison_dir / relpath(entity["legacy_knowledge_bridger"]["copied"], comparison_dir)
            word_original = prepare_word_image(original, cache_dir)
            word_mmkgc = prepare_word_image(mmkgc, cache_dir)
            word_legacy = prepare_word_image(legacy, cache_dir)
            mmkgc_detail = (
                f"候选 {entity['mmkgc']['candidate_id']}\n"
                f"总分 {format_score(entity['mmkgc']['score'])}\n"
                f"CLIP {format_score(entity['mmkgc'].get('clip_score'))}\n"
                f"BLIP {format_score(entity['mmkgc'].get('blip_itm_score'))}\n"
                f"去Judge {format_score(entity['mmkgc'].get('clip_blip_contribution'))}\n"
                f"CLIP+BLIP归一化 {format_score(entity['mmkgc'].get('clip_blip_score'))}"
            )
            legacy_detail = (
                f"候选 id_{entity['legacy_knowledge_bridger']['candidate_index']}.jpg\n"
                f"CLIP {format_score(entity['legacy_knowledge_bridger'].get('clip_score'))}\n"
                f"BLIP {format_score(entity['legacy_knowledge_bridger'].get('blip_itm_score'))}\n"
                f"去Judge {format_score(entity['legacy_knowledge_bridger'].get('clip_blip_contribution'))}\n"
                f"CLIP+BLIP归一化 "
                f"{format_score(entity['legacy_knowledge_bridger'].get('clip_blip_score'))}"
            )

            md_parts.extend([
                f'<a id="entity-{idx:03d}"></a>',
                f"## {idx + 1}. {md_escape(label)}",
                "",
                f"- **实体 ID**：`{key}`",
                f"- **类别**：{md_escape(entity.get('category', row.get('category', '')))}",
                "",
                "### 图片对比",
                "",
                "| 原图 | MMKGC 最高分图片 | legacy 统一重评分最高分图片 |",
                "|---|---|---|",
                f"| ![原图]({relpath(original, comparison_dir)}) | ![MMKGC 最优图]({relpath(mmkgc, comparison_dir)})<br>候选 `{md_escape(entity['mmkgc']['candidate_id'])}`<br>总分：`{format_score(entity['mmkgc']['score'])}`<br>CLIP：`{format_score(entity['mmkgc'].get('clip_score'))}`<br>BLIP：`{format_score(entity['mmkgc'].get('blip_itm_score'))}`<br>去 Judge：`{format_score(entity['mmkgc'].get('clip_blip_contribution'))}`<br>CLIP+BLIP 归一化：`{format_score(entity['mmkgc'].get('clip_blip_score'))}` | ![legacy 重评分最优图]({relpath(legacy, comparison_dir)})<br>候选 `id_{entity['legacy_knowledge_bridger']['candidate_index']}.jpg`<br>CLIP：`{format_score(entity['legacy_knowledge_bridger'].get('clip_score'))}`<br>BLIP：`{format_score(entity['legacy_knowledge_bridger'].get('blip_itm_score'))}`<br>去 Judge：`{format_score(entity['legacy_knowledge_bridger'].get('clip_blip_contribution'))}`<br>CLIP+BLIP 归一化：`{format_score(entity['legacy_knowledge_bridger'].get('clip_blip_score'))}` |",
                "",
                "### 文本描述",
                "",
                description,
                "",
                "### MMKGC 图片生成 Prompt",
                "",
                f"候选 `{md_escape(prompts['mmkgc_candidate'])}`：",
                "",
                f"```text\n{prompts['mmkgc']}\n```",
                "",
                "### legacy Knowledge-Bridger 图片生成 Prompt",
                "",
                f"候选 `{md_escape(prompts['legacy_candidate'])}`：",
                "",
                f"```text\n{prompts['legacy']}\n```",
                "",
                "---",
                "",
            ])

            if idx > 0:
                document.add_page_break()
            document.add_heading(f"{idx + 1}. {label}", level=1)
            document.add_paragraph(f"实体 ID：{key}    类别：{entity.get('category', row.get('category', ''))}")
            table = document.add_table(rows=1, cols=3)
            table.autofit = True
            cells = table.rows[0].cells
            add_image_cell(cells[0], "原图", word_original, "")
            add_image_cell(cells[1], "MMKGC 最高分图片", word_mmkgc, mmkgc_detail)
            add_image_cell(cells[2], "legacy 统一重评分最高分图片", word_legacy, legacy_detail)
            score_note = document.add_paragraph(
                f"分数对比：MMKGC 总分 {format_score(entity['mmkgc']['score'])}，"
                f"CLIP {format_score(entity['mmkgc'].get('clip_score'))}，"
                f"BLIP {format_score(entity['mmkgc'].get('blip_itm_score'))}，"
                f"去 Judge 分数 {format_score(entity['mmkgc'].get('clip_blip_contribution'))}，"
                f"CLIP+BLIP 归一化 {format_score(entity['mmkgc'].get('clip_blip_score'))}；"
                f"legacy CLIP {format_score(entity['legacy_knowledge_bridger'].get('clip_score'))}，"
                f"BLIP {format_score(entity['legacy_knowledge_bridger'].get('blip_itm_score'))}，"
                f"去 Judge 分数 "
                f"{format_score(entity['legacy_knowledge_bridger'].get('clip_blip_contribution'))}，"
                f"CLIP+BLIP 归一化 "
                f"{format_score(entity['legacy_knowledge_bridger'].get('clip_blip_score'))}。"
            )
            for run in score_note.runs:
                run.font.size = Pt(8)
            document.add_heading("文本描述", level=2)
            document.add_paragraph(description)
            document.add_heading("MMKGC 图片生成 Prompt", level=2)
            document.add_paragraph(f"候选：{prompts['mmkgc_candidate']}")
            mmkgc_prompt_paragraph = document.add_paragraph(prompts["mmkgc"])
            for run in mmkgc_prompt_paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            document.add_heading("legacy Knowledge-Bridger 图片生成 Prompt", level=2)
            document.add_paragraph(f"候选：{prompts['legacy_candidate']}")
            legacy_prompt_paragraph = document.add_paragraph(prompts["legacy"])
            for run in legacy_prompt_paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(8)

    markdown_path = comparison_dir / "comparison_report.md"
    markdown_path.write_text("\n".join(md_parts), encoding="utf-8")
    docx_path = comparison_dir / "comparison_report.docx"
    document.save(docx_path)
    print(f"markdown={markdown_path}")
    print(f"docx={docx_path}")
    print(f"entities={len(manifest['entities'])}")


if __name__ == "__main__":
    main()
