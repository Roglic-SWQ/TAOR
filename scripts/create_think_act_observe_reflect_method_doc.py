from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "think-act-observe-reflect-method.zh.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_cell_text(cell, text: str, *, bold=False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        set_cell_text(header.cells[index], value, bold=True, color="FFFFFF")
        set_cell_shading(header.cells[index], "1F4E79")
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            set_cell_text(row.cells[index], value)
            if len(table.rows) % 2 == 0:
                set_cell_shading(row.cells[index], "F2F6FA")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(text)


def add_code_block(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.add_run(text)


def add_formula(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Formula")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(text)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in (
        ("Title", 20, "1F1F1F"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 12.5, "2F5597"),
        ("Heading 3", 11, "404040"),
    ):
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    if "Code Block" not in styles:
        style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Code Block"]
    style.font.name = "Consolas"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    style.font.size = Pt(9)
    style.font.color.rgb = RGBColor.from_string("333333")
    style.paragraph_format.left_indent = Cm(0.6)
    style.paragraph_format.right_indent = Cm(0.6)
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.05

    if "Formula" not in styles:
        style = styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Formula"]
    style.font.name = "Cambria Math"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    style.font.size = Pt(10.5)
    style.font.italic = True
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(8)


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("MMKGC | Think–Act–Observe–Reflect 方法草案")
    run.font.name = "等线"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("7F7F7F")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("方法章节草稿")
    run.font.name = "等线"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("7F7F7F")


def build_document() -> Document:
    document = Document()
    configure_document(document)
    add_header_footer(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("知识图谱约束的 Think–Act–Observe–Reflect 多模态补全方法")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("论文方法章节草稿（面向 MMKGC 图像补全任务）")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("666666")

    document.add_paragraph()
    abstract = document.add_paragraph()
    abstract.add_run("摘要式概述：").bold = True
    abstract.add_run(
        "本文提出一种知识图谱约束的 Think–Act–Observe–Reflect 多模态补全框架，"
        "用于在已有实体模态和局部知识图谱证据不完整的情况下生成缺失的文本或图像模态。"
        "框架首先对局部关系进行高召回筛选，并构建带有事实可信度和证据引用的实体知识模型；"
        "随后通过关系重排序确定生成重点，将关系计划编译为受约束的 prompt；"
        "生成结果再由 KG-aware 的观察模块进行事实一致性、关系一致性和内容覆盖评估；"
        "最后根据违反项和遗漏项构造有限的反思反馈，用于安全重生成。"
        "该设计将知识图谱从静态输入转化为贯穿规划、生成和评估的可追溯约束。"
    )

    document.add_heading("1. 问题定义", level=1)
    document.add_paragraph(
        "给定目标实体 e 的已有模态 x_e、局部知识图谱 G_e，以及需要补全的缺失模态 m，"
        "目标是利用已有模态和知识图谱中的结构化事实与关系证据，生成缺失模态的候选 y，"
        "使其与实体身份、已有模态语义和图谱事实保持一致。"
        "本文重点处理 image completion，也保留 text completion 作为统一接口。"
    )
    add_formula(document, "G_e = (V_e, E_e),    y_m* = argmax_y  Consistency(y | x_e, G_e)  subject to  Valid(y, m)")
    document.add_paragraph(
        "其中，V_e 和 E_e 分别表示目标实体局部邻域中的实体和关系边；"
        "Valid(y, m) 表示候选满足目标模态、格式和生成器约束；"
        "Consistency(y, G_e) 表示候选与知识图谱证据、实体画像和生成约束的一致程度。"
    )

    document.add_heading("2. Think–Act–Observe–Reflect 总体框架", level=1)
    document.add_paragraph(
        "现有 MMKGC 五阶段主链路可以在更高层抽象为四个认知环节。"
        "Think 对应关系粗筛、知识建模、可信度评估和关系重排序；"
        "Act 对应受关系计划控制的 prompt 生成和图像生成；"
        "Observe 对应 KG-aware 候选一致性排序；"
        "Reflect 对应根据观察结果进行的一次受约束 safer regeneration。"
        "该抽象不改变五个阶段的工程职责，而是为方法提供统一的闭环解释。"
    )
    add_table(
        document,
        ["闭环环节", "对应阶段", "核心问题", "主要输出"],
        [
            ["Think", "Stage 0–2", "哪些证据可靠，哪些关系值得驱动生成？", "知识模型、事实可信度、关系计划"],
            ["Act", "Stage 3", "如何把关系计划转化为生成行为？", "prompt、文本候选或图像候选"],
            ["Observe", "Stage 4", "候选是否符合实体、事实和关系约束？", "质量分数、违反项、覆盖信号"],
            ["Reflect", "反馈到 Stage 3，必要时反馈到 Stage 2", "如何修复违反和遗漏？", "结构化修正反馈、重新生成结果"],
        ],
        widths=[2.2, 2.4, 6.0, 5.0],
    )
    add_formula(document, "K = Think(G_e, x_e),    P = Act(K, m),    y = Generate(P),    O = Observe(y, K),    P' = Reflect(P, O, K)")

    document.add_heading("3. Think：知识建模与关系重排序", level=1)
    document.add_paragraph(
        "Think 阶段不要求生成模型自由地产生事实，而是首先建立一个可追溯、可分级使用的知识模型。"
        "其输入为局部图谱证据和已有模态，输出为实体画像、关系事实、可信度标签、模态约束和关系生成计划。"
    )

    document.add_heading("3.1 局部关系筛选", level=2)
    document.add_paragraph(
        "对于目标实体，从一跳邻域及有限路径上下文中获得候选关系集合："
    )
    add_formula(document, "R_e = {r_1, r_2, ..., r_n}")
    document.add_paragraph(
        "系统首先按关系和方向聚合邻居，保留频次、邻居实体、证据三元组和路径上下文，"
        "并过滤 ID、URL、归档和其他低生成价值的元数据关系。该步骤的目标是高召回地压缩知识邻域，"
        "而不是在此处完成最终关系选择。"
    )

    document.add_heading("3.2 知识建模与事实可信度", level=2)
    document.add_paragraph(
        "每个候选关系被转换为带证据引用的事实对象："
    )
    add_formula(document, "f_i = (e, r_i, o_i, d_i, E_i, c_i)")
    document.add_paragraph(
        "其中 r_i 为关系，o_i 为邻居实体，d_i 为关系方向，E_i 为支持事实的证据集合，"
        "c_i 为最终可信度。可信度由确定性证据先验和有限的语义修正共同决定："
    )
    add_formula(document, "c_i = clip(c_i^prior + Δ_i, 0, 1),    Δ_i ∈ [-δ, δ]")
    document.add_paragraph(
        "语义模型只能针对已有 fact_id 返回有限的修正量、风险等级和理由，不能新增事实、修改证据或直接指定生成权限。"
        "根据最终可信度和风险，事实被编译为以下四类使用策略："
    )
    add_table(
        document,
        ["事实类别", "生成使用方式", "含义"],
        [
            ["direct_generation", "must_include", "有直接证据支持，可作为核心生成内容"],
            ["supporting_only", "should_include", "可以辅助生成，但不应替代实体身份"],
            ["context_only", "context_only", "只提供背景，不作为确定事实表达"],
            ["excluded", "不进入正向约束", "风险过高或证据不足，禁止正向生成"],
        ],
        widths=[4.0, 4.0, 7.6],
    )

    document.add_heading("3.3 关系重排序与生成计划", level=2)
    document.add_paragraph(
        "为了使有限的候选预算集中于最有价值的关系，系统综合模态适配性、约束匹配、证据强度、生成效用和风险惩罚进行关系评分："
    )
    add_formula(document, "S(r_i) = λ_1 M_i + λ_2 A_i + λ_3 E_i + λ_4 U_i - λ_5 Q_i")
    document.add_paragraph(
        "其中 M_i 表示关系与目标模态的适配性，A_i 表示关系与生成约束的匹配度，"
        "E_i 表示证据强度，U_i 表示关系作为生成依据的效用，Q_i 表示风险惩罚。"
        "关系按分数和硬排除条件划分为 primary、secondary 和 context 三种角色，并进一步分配候选配额。"
    )
    add_code_block(
        document,
        '{\n'
        '  "primary_relations": ["instance of", "creator"],\n'
        '  "secondary_relations": ["genre"],\n'
        '  "context_relations": ["country"],\n'
        '  "must_include": ["entity type", "creator"],\n'
        '  "should_include": ["genre"]\n'
        '}',
    )

    document.add_heading("4. Act：关系条件化 Prompt 与图像生成", level=1)
    document.add_paragraph(
        "Act 阶段将 Think 阶段产生的关系计划、实体画像和事实约束编译为可执行的生成指令。"
        "与直接将原始 KG 拼接到 prompt 不同，本文只将允许使用的事实和目标模态相关的关系传递给生成器，"
        "同时明确限制模型不得补充未被证据支持的精确细节。"
    )

    document.add_heading("4.1 Prompt 规划", level=2)
    document.add_paragraph(
        "对目标模态 m，prompt 由实体身份、主要关系、支持事实、视觉或文本场景线索以及不确定性边界共同组成："
    )
    add_formula(document, "P_m = PromptPlanner(K, R_primary, R_secondary, C_m)")
    add_bullet(document, "实体身份：实体类型、核心身份和语义领域。")
    add_bullet(document, "关系重点：当前候选对应的 primary relation 及其支持关系。")
    add_bullet(document, "证据约束：must_include、should_include 和 context_only 的分层内容。")
    add_bullet(document, "安全边界：禁止把不确定事实转换为确定的视觉或文本细节。")

    document.add_heading("4.2 多候选图像生成", level=2)
    document.add_paragraph(
        "为了避免单一 prompt 对生成结果的偶然性，系统围绕不同生成角度产生多个候选："
        "identity anchor、evidence-grounded description、relation focus、attribute/context 和 semantic bridge。"
        "每个候选都绑定其 primary relation、supporting relations、证据三元组和使用的约束。"
    )
    add_formula(document, "I_j = G(P_j, z_j),    j = 1, ..., B")
    document.add_paragraph(
        "其中 B 为候选预算，P_j 为关系条件化 prompt，z_j 表示采样随机性或生成条件。"
        "对于图像补全，Qwen prompt planner 负责生成图像 prompt，SDXL 负责生成实际 PNG artifact；"
        "对于文本补全，则由文本生成器直接生成候选文本。"
    )

    document.add_heading("5. Observe：生成结果的知识一致性观察", level=1)
    document.add_paragraph(
        "Observe 阶段不重新改变知识模型，而是检查候选是否符合 Think 阶段确定的证据边界和关系计划。"
        "它将生成质量、语义对齐、结构化知识对齐和 Judge 诊断组合为候选观察结果。"
    )
    add_formula(document, "O(y, K) = (q(y), a(y, K), g(y, K), V(y, K), M(y, K))")
    document.add_paragraph(
        "q 表示生成质量，a 表示语义和模态对齐，g 表示与知识图谱事实的结构化对齐，"
        "V 表示违反项，M 表示遗漏项。"
    )

    document.add_heading("5.1 事实与关系违反", level=2)
    document.add_paragraph("观察模块重点识别以下类型的问题：")
    add_bullet(document, "实体身份不一致，例如实体类型或核心身份发生偏移。")
    add_bullet(document, "候选引入了原始 KG 和允许事实中不存在的关系或属性。")
    add_bullet(document, "候选内容与 KG 支持的关系方向、邻居实体或事实发生冲突。")
    add_bullet(document, "候选使用了 excluded 或 high-risk 信息。")
    add_bullet(document, "图像或文本没有体现当前候选的主要关系。")

    document.add_heading("5.2 关系覆盖与内容遗漏", level=2)
    document.add_paragraph(
        "每个候选保留 expected_relation_coverage，用于记录该候选理论上应表达的关系。"
        "观察模块将其与候选实际表达出的事实或视觉语义进行比较，形成覆盖信号。"
        "对于文本候选，可以通过结构化三元组抽取与 KG 三元组对齐计算 Precision、Recall 和 F1；"
        "对于图像候选，则使用图文相似度、视觉判别器和 KG-aware Judge 的联合信号评估关系表达。"
    )
    add_formula(document, "Precision = |T_pred ∩ T_KG| / |T_pred|,    Recall = |T_pred ∩ T_KG| / |T_KG|")
    add_formula(document, "F1 = 2 · Precision · Recall / (Precision + Recall)")
    document.add_paragraph(
        "在论文表述中，观察模块不应只报告一个总分，还应将低 Recall 解释为潜在遗漏，将低 Precision 解释为潜在无依据内容。"
        "当前工程实现已经保留候选关系覆盖、结构化 KG 对齐、hard/soft violation 和 Judge evidence 等诊断字段，"
        "这些字段共同构成 Observe 的输出。"
    )

    document.add_heading("6. Reflect：约束保持的反馈与安全重生成", level=1)
    document.add_paragraph(
        "Reflect 阶段根据 Observe 的结果生成修正反馈，但不重新发明事实，也不允许模型绕过知识模型。"
        "其作用是将观察结果转换为下一次生成可执行的限制。"
    )
    add_formula(document, "Δ = Reflect(O, K, P),    P' = UpdatePrompt(P, Δ)")
    add_table(
        document,
        ["观察问题", "反馈动作", "是否适合重生成"],
        [
            ["unsupported claim / fact", "删除不受支持的内容", "是"],
            ["relation or identity conflict", "保持实体身份并修正关系表达", "是；严重时重新选择关系"],
            ["missing core evidence", "补充已有 KG 支持的核心事实", "是"],
            ["low relation coverage", "增强 selected relation 的表达", "是"],
            ["generator/schema/service failure", "记录系统失败，不把它伪装成内容问题", "否；需由系统处理"],
        ],
        widths=[5.0, 7.0, 4.0],
    )
    document.add_paragraph(
        "反思反馈可以使用如下结构化表示："
    )
    add_code_block(
        document,
        '{\n'
        '  "repair_scope": "stage3",\n'
        '  "required_fact_refs": ["fact:creator"],\n'
        '  "remove_fact_refs": [],\n'
        '  "remove_relations": ["unsupported location"],\n'
        '  "missing_relations": ["creator"],\n'
        '  "repair_actions": [\n'
        '    "include graph-supported creator evidence",\n'
        '    "remove unsupported location claims"\n'
        '  ]\n'
        '}',
    )
    document.add_paragraph(
        "本文采用有限反思策略。若候选没有有效结果，或观察结果包含可由生成器修复的严重违反项，"
        "系统构造 safer regeneration feedback 并重新生成一次候选；否则直接保留当前排序结果。"
        "这种设计保证反思过程可审计、成本可控，也避免生成系统陷入无界循环。"
    )

    document.add_heading("7. 统一推理算法", level=1)
    add_code_block(
        document,
        "Algorithm 1: KG-Grounded Think–Act–Observe–Reflect Completion\n"
        "Input: entity e, local KG G_e, existing modality x_e, target modality m\n"
        "1:  C <- CoarseRelationFiltering(G_e, m)\n"
        "2:  K <- KnowledgeModeling(C, x_e)\n"
        "3:  K <- CredibilityAssessmentAndConstraintCompilation(K)\n"
        "4:  R <- RelationReranking(K, m)\n"
        "5:  P <- PromptPlanning(K, R, m)\n"
        "6:  Y <- GenerateCandidates(P)\n"
        "7:  O <- ObserveCandidates(Y, K, R)\n"
        "8:  if O contains repairable violations or unacceptable omissions then\n"
        "9:      Δ <- Reflect(O, K, R)\n"
        "10:     P' <- UpdatePrompt(P, Δ)\n"
        "11:     Y' <- GenerateCandidates(P')\n"
        "12:     O' <- ObserveCandidates(Y', K, R)\n"
        "13:     return BestValidCandidate(Y ∪ Y')\n"
        "14: else\n"
        "15:     return BestValidCandidate(Y)\n"
        "16: end if",
    )

    document.add_heading("8. 与 MMKGC 当前实现的对应关系", level=1)
    document.add_paragraph(
        "该方法抽象与当前代码的对应关系如下。论文中应明确说明：Think–Act–Observe–Reflect 是对现有五阶段主链路的统一方法解释，"
        "而不是声称已经实现了无限多轮自主反思。"
    )
    add_table(
        document,
        ["方法概念", "当前实现", "说明"],
        [
            ["Think", "Stage 0–2", "coarse filtering、knowledge modeling、credibility/constraint compilation、relation selection"],
            ["Act", "Stage 3", "关系条件化候选生成；图像路径为 Qwen prompt planner + SDXL"],
            ["Observe", "Stage 4", "hard gate、文本结构化 KG 对齐、CLIP/BLIP、Visual Judge、KG-aware Judge"],
            ["Reflect", "Stage 3 feedback", "当前为一次 safer regeneration；反馈来自 Stage 4 的可修复 violations"],
        ],
        widths=[3.0, 4.0, 10.0],
    )

    document.add_heading("9. 方法边界与论文表述建议", level=1)
    add_bullet(document, "将该方法描述为知识图谱约束的有限闭环，而不是通用自主智能体。")
    add_bullet(document, "强调关系计划是 Think 与 Act 之间的中间表示，使知识证据可追踪到 prompt 和候选。")
    add_bullet(document, "将 hard violation、soft violation、结构化对齐和 Judge evidence 统一作为 Observe 信号。")
    add_bullet(document, "明确 Reflect 的修正范围：删除不支持内容、补充已有证据、强化已选关系，不创造新事实。")
    add_bullet(document, "如果视觉关系抽取尚未完全稳定，应将图像侧关系遗漏描述为联合观察信号，而不要声称有完备视觉场景图推理。")

    document.add_heading("10. 论文方法章节结构建议", level=1)
    for item in (
        "3.1 Problem Definition",
        "3.2 Overview of the Think–Act–Observe–Reflect Framework",
        "3.3 Think: Knowledge Modeling and Relation Reranking",
        "3.4 Act: Relation-conditioned Prompt Planning and Image Generation",
        "3.5 Observe: KG-aware Multimodal Consistency Evaluation",
        "3.6 Reflect: Constraint-preserving Regeneration",
        "3.7 Candidate Selection and Inference Procedure",
    ):
        add_number(document, item)

    document.add_heading("11. 可概括的论文贡献", level=1)
    add_number(document, "提出一个以知识图谱为证据边界的 Think–Act–Observe–Reflect 多模态补全框架。")
    add_number(document, "将关系重排序结果显式编译为 prompt 生成约束，降低无依据内容和关系偏移。")
    add_number(document, "通过生成后观察与有限反思反馈，统一处理事实违反、关系不一致和内容遗漏。")

    document.add_paragraph()
    note = document.add_paragraph()
    note.add_run("说明：").bold = True
    note.add_run(
        "本文档是根据当前 MMKGC 五阶段主链路整理的方法章节草稿。"
        "其中 Observe 和 Reflect 的表述保留了当前实现边界；若后续新增独立的关系覆盖检测或多轮反思控制器，"
        "可在第 5、6 节中继续扩展。"
    )
    return document


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)
